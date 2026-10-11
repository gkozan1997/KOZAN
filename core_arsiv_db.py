import os
import sys
import time
import sqlite3
import json
import tempfile
import threading
from datetime import datetime, timezone, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def _get_cloud_module():
    try:
        import core_arsiv_cloud
        return core_arsiv_cloud
    except ImportError:
        try:
            import core_cloud_sync
            return core_cloud_sync
        except ImportError:
            return None

TZ_ISTANBUL = timezone(timedelta(hours=3))

IS_CLOUD = (
    os.environ.get('VERCEL') == '1' or
    os.environ.get('AWS_LAMBDA_FUNCTION_NAME') is not None or
    os.environ.get('NOW_REGION') is not None
)

def _sync_cloud_record(target_func, *args, **kwargs):
    # Vercel serverless ortaminda thread'ler response donunce donduruldugu icin dogrudan ve senkron kaydedilmeli!
    if IS_CLOUD:
        try:
            target_func(*args, **kwargs)
        except Exception as e:
            print(f"[CloudArchiveSync] Senkron yazma hatasi: {e}")
    else:
        def _worker():
            try:
                target_func(*args, **kwargs)
            except Exception as e:
                print(f"[CloudArchiveWorker] Hata: {e}")
        t = threading.Thread(target=_worker, daemon=True)
        t.start()

_LAST_CLOUD_SYNC_TS = 0

def _ensure_cloud_synced():
    global _LAST_CLOUD_SYNC_TS
    now = time.time()
    # Eger buluttaysak ve son 5 saniyede senkron yapilmadiysa buluttan verileri cek
    if (IS_CLOUD or not _has_any_records()) and (now - _LAST_CLOUD_SYNC_TS > 5):
        try:
            cmod = _get_cloud_module()
            if cmod and hasattr(cmod, 'sync_db_with_cloud_archive'):
                cmod.sync_db_with_cloud_archive()
                _LAST_CLOUD_SYNC_TS = now
        except Exception as e:
            print(f"[CloudAutoSync] Otomatik senkron hatasi: {e}")

def _has_any_records():
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM barkod_kayitlari WHERE durum = 'Gecerli'")
        row = cursor.fetchone()
        conn.close()
        return (row[0] > 0) if row else False
    except Exception:
        return False

def get_now_istanbul():
    return datetime.now(TZ_ISTANBUL)

def get_base_data_dir():
    if IS_CLOUD:
        base = os.path.join(tempfile.gettempdir(), 'barkod_arsiv_data')
    else:
        # Ortak PROJLERİM/barkod_arsiv/data klasörünü kontrol et (Masaüstü senkronizasyonu)
        desktop_candidates = [
            os.path.expanduser('~/OneDrive/Masaüstü/PROJLERİM/barkod_arsiv/data'),
            os.path.expanduser('~/Desktop/PROJLERİM/barkod_arsiv/data'),
            'C:/Users/ugurk/OneDrive/Masaüstü/PROJLERİM/barkod_arsiv/data',
        ]
        found_data = None
        for candidate in desktop_candidates:
            if os.path.exists(candidate):
                found_data = candidate
                break
        if found_data:
            base = found_data
        else:
            base_dir = os.path.dirname(os.path.abspath(__file__))
            base = os.path.join(base_dir, 'data')
    os.makedirs(base, exist_ok=True)
    return base

def get_db_path():
    data_dir = get_base_data_dir()
    return os.path.join(data_dir, 'barkod_arsiv.db')

def get_backup_dir():
    data_dir = get_base_data_dir()
    backup_dir = os.path.join(data_dir, 'arsiv_gunluk')
    os.makedirs(backup_dir, exist_ok=True)
    return backup_dir

def get_connection():
    db_path = get_db_path()
    conn = sqlite3.connect(db_path, timeout=15)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS barkod_kayitlari (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kargo_barkodu TEXT,
            siparis_no TEXT,
            musteri_adi TEXT,
            urun_adi TEXT,
            barkod_seri_no TEXT NOT NULL,
            pazaryeri TEXT,
            kayit_tarihi TEXT NOT NULL,
            kayit_gunu TEXT NOT NULL,
            okutma_yontemi TEXT DEFAULT 'El Terminali',
            notlar TEXT,
            durum TEXT DEFAULT 'Gecerli'
        )
    ''')
    # Kolon var mı kontrolü (Migration)
    cursor.execute("PRAGMA table_info(barkod_kayitlari)")
    cols = [r[1] for r in cursor.fetchall()]
    if 'kargo_barkodu' not in cols:
        cursor.execute("ALTER TABLE barkod_kayitlari ADD COLUMN kargo_barkodu TEXT")

    cursor.execute('CREATE INDEX IF NOT EXISTS idx_kargo_barkod ON barkod_kayitlari(kargo_barkodu)')
    cursor.execute('CREATE INDEX IF NOT EXISTS idx_barkod ON barkod_kayitlari(barkod_seri_no)')
    cursor.execute('CREATE INDEX IF NOT EXISTS idx_siparis ON barkod_kayitlari(siparis_no)')
    cursor.execute('CREATE INDEX IF NOT EXISTS idx_gun ON barkod_kayitlari(kayit_gunu)')
    conn.commit()
    conn.close()

def check_duplicate(barkod_seri_no):
    if not barkod_seri_no:
        return False, None
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute('''
            SELECT id, kargo_barkodu, siparis_no, musteri_adi, urun_adi, barkod_seri_no, kayit_tarihi, pazaryeri
            FROM barkod_kayitlari
            WHERE barkod_seri_no = ? AND durum = 'Gecerli'
            ORDER BY id DESC LIMIT 1
        ''', (barkod_seri_no.strip(),))
        row = cursor.fetchone()
        conn.close()
        if row:
            return True, dict(row)
    except Exception as e:
        print(f"Mükerrer kontrol hatası: {e}")
    return False, None

def add_barcode_record(kargo_barkodu='', siparis_no='', musteri_adi='', urun_adi='', barkod_seri_no='', pazaryeri='Bilinmiyor', okutma_yontemi='El Terminali', notlar=''):
    now = get_now_istanbul()
    kayit_tarihi = now.strftime('%Y-%m-%d %H:%M:%S')
    kayit_gunu = now.strftime('%Y-%m-%d')
    barkod_temiz = str(barkod_seri_no).strip()
    kargo_temiz = str(kargo_barkodu or '').strip()

    is_dup, dup_info = check_duplicate(barkod_temiz)

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO barkod_kayitlari (
            kargo_barkodu, siparis_no, musteri_adi, urun_adi, barkod_seri_no,
            pazaryeri, kayit_tarihi, kayit_gunu, okutma_yontemi, notlar, durum
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Gecerli')
    ''', (
        kargo_temiz,
        (siparis_no or '').strip(),
        (musteri_adi or '').strip(),
        (urun_adi or '').strip(),
        barkod_temiz,
        (pazaryeri or 'Bilinmiyor').strip(),
        kayit_tarihi,
        kayit_gunu,
        okutma_yontemi,
        (notlar or '').strip()
    ))
    record_id = cursor.lastrowid
    conn.commit()
    conn.close()

    record_data = {
        'id': record_id,
        'kargo_barkodu': kargo_temiz,
        'siparis_no': siparis_no or '',
        'musteri_adi': musteri_adi or '',
        'urun_adi': urun_adi or '',
        'barkod_seri_no': barkod_temiz,
        'pazaryeri': pazaryeri or '',
        'kayit_tarihi': kayit_tarihi,
        'kayit_gunu': kayit_gunu,
        'okutma_yontemi': okutma_yontemi,
        'notlar': notlar or '',
        'is_duplicate': is_dup
    }

    try:
        append_to_daily_json(kayit_gunu, record_data)
    except Exception as e:
        print(f"JSON yedekleme hatası: {e}")

    try:
        cmod = _get_cloud_module()
        if cmod and hasattr(cmod, 'append_archive_record_to_cloud'):
            _sync_cloud_record(cmod.append_archive_record_to_cloud, record_data)
    except Exception as e:
        print(f"[CloudArchive] Async append hatasi: {e}")

    return record_data, is_dup, dup_info

def append_to_daily_json(gun_str, record_data):
    backup_dir = get_backup_dir()
    json_path = os.path.join(backup_dir, f'{gun_str}.json')
    records = []
    if os.path.exists(json_path):
        try:
            with open(json_path, 'r', encoding='utf-8') as f:
                records = json.load(f)
        except Exception:
            records = []
    records.append(record_data)
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

def remove_from_daily_json(gun_str, record_id):
    if not gun_str:
        return
    backup_dir = get_backup_dir()
    json_path = os.path.join(backup_dir, f'{gun_str}.json')
    if os.path.exists(json_path):
        try:
            with open(json_path, 'r', encoding='utf-8') as f:
                records = json.load(f)
            records = [r for r in records if r.get('id') != record_id]
            with open(json_path, 'w', encoding='utf-8') as f:
                json.dump(records, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"JSON silme hatasi: {e}")

def update_in_daily_json(gun_str, record_id, fields):
    if not gun_str:
        return
    backup_dir = get_backup_dir()
    json_path = os.path.join(backup_dir, f'{gun_str}.json')
    if os.path.exists(json_path):
        try:
            with open(json_path, 'r', encoding='utf-8') as f:
                records = json.load(f)
            for r in records:
                if r.get('id') == record_id:
                    for k, v in fields.items():
                        r[k] = v
            with open(json_path, 'w', encoding='utf-8') as f:
                json.dump(records, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"JSON guncelleme hatasi: {e}")

def delete_record_by_id(record_id):
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM barkod_kayitlari WHERE id = ?", (record_id,))
        row = cursor.fetchone()
        if not row:
            conn.close()
            return False, "Kayıt bulunamadı"

        rec = dict(row)
        cursor.execute("DELETE FROM barkod_kayitlari WHERE id = ?", (record_id,))
        conn.commit()
        conn.close()

        remove_from_daily_json(rec.get('kayit_gunu'), record_id)
        try:
            cmod = _get_cloud_module()
            if cmod and hasattr(cmod, 'delete_archive_record_from_cloud'):
                _sync_cloud_record(cmod.delete_archive_record_from_cloud, record_id=record_id)
        except Exception as e:
            print(f"[CloudArchive] Async delete hatasi: {e}")
        return True, "Kayıt başarıyla silindi"
    except Exception as e:
        return False, str(e)

def delete_record_by_barcode(barkod_seri_no, kargo_barkodu='', siparis_no=''):
    try:
        conn = get_connection()
        cursor = conn.cursor()
        sql = "SELECT id, kayit_gunu FROM barkod_kayitlari WHERE barkod_seri_no = ?"
        params = [str(barkod_seri_no).strip()]
        if kargo_barkodu:
            sql += " AND kargo_barkodu = ?"
            params.append(str(kargo_barkodu).strip())
        elif siparis_no:
            sql += " AND siparis_no = ?"
            params.append(str(siparis_no).strip())

        sql += " ORDER BY id DESC LIMIT 1"
        cursor.execute(sql, params)
        row = cursor.fetchone()
        if not row:
            cursor.execute("SELECT id, kayit_gunu FROM barkod_kayitlari WHERE barkod_seri_no = ? ORDER BY id DESC LIMIT 1", (str(barkod_seri_no).strip(),))
            row = cursor.fetchone()

        if row:
            rec_id = row['id']
            rec_gun = row['kayit_gunu']
            cursor.execute("DELETE FROM barkod_kayitlari WHERE id = ?", (rec_id,))
            conn.commit()
            conn.close()
            remove_from_daily_json(rec_gun, rec_id)
            try:
                cmod = _get_cloud_module()
                if cmod and hasattr(cmod, 'delete_archive_record_from_cloud'):
                    _sync_cloud_record(cmod.delete_archive_record_from_cloud, record_id=rec_id, barkod_seri_no=barkod_seri_no, kargo_barkodu=kargo_barkodu)
            except Exception as e:
                print(f"[CloudArchive] Async delete hatasi: {e}")
            return True, "Kayıt başarıyla silindi"

        conn.close()
        return True, "Kayıt silindi"
    except Exception as e:
        return False, str(e)

def update_record_by_id(record_id, fields):
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM barkod_kayitlari WHERE id = ?", (record_id,))
        row = cursor.fetchone()
        if not row:
            conn.close()
            return False, "Kayıt bulunamadı"

        allowed = ['barkod_seri_no', 'kargo_barkodu', 'siparis_no', 'musteri_adi', 'urun_adi', 'pazaryeri', 'notlar', 'durum']
        sets = []
        params = []
        for k in allowed:
            if k in fields:
                sets.append(f"{k} = ?")
                params.append(str(fields[k]).strip())

        if not sets:
            conn.close()
            return True, "Değişiklik yok"

        params.append(record_id)
        cursor.execute(f"UPDATE barkod_kayitlari SET {', '.join(sets)} WHERE id = ?", params)
        conn.commit()
        conn.close()

        update_in_daily_json(row['kayit_gunu'], record_id, fields)
        try:
            cmod = _get_cloud_module()
            if cmod and hasattr(cmod, 'update_archive_record_in_cloud'):
                _sync_cloud_record(cmod.update_archive_record_in_cloud, record_id=record_id, fields=fields)
        except Exception as e:
            print(f"[CloudArchive] Async update hatasi: {e}")
        return True, "Kayıt başarıyla güncellendi"
    except Exception as e:
        return False, str(e)

def update_record_by_barcode(old_barcode, new_barcode, kargo_barkodu='', siparis_no=''):
    try:
        conn = get_connection()
        cursor = conn.cursor()
        sql = "SELECT id, kayit_gunu FROM barkod_kayitlari WHERE barkod_seri_no = ?"
        params = [str(old_barcode).strip()]
        if kargo_barkodu:
            sql += " AND kargo_barkodu = ?"
            params.append(str(kargo_barkodu).strip())
        elif siparis_no:
            sql += " AND siparis_no = ?"
            params.append(str(siparis_no).strip())

        sql += " ORDER BY id DESC LIMIT 1"
        cursor.execute(sql, params)
        row = cursor.fetchone()
        if not row:
            cursor.execute("SELECT id, kayit_gunu FROM barkod_kayitlari WHERE barkod_seri_no = ? ORDER BY id DESC LIMIT 1", (str(old_barcode).strip(),))
            row = cursor.fetchone()

        if row:
            rec_id = row['id']
            clean_new = str(new_barcode).strip()
            cursor.execute("UPDATE barkod_kayitlari SET barkod_seri_no = ? WHERE id = ?", (clean_new, rec_id))
            conn.commit()
            conn.close()
            update_in_daily_json(row['kayit_gunu'], rec_id, {'barkod_seri_no': clean_new})
            try:
                cmod = _get_cloud_module()
                if cmod and hasattr(cmod, 'update_archive_record_in_cloud'):
                    _sync_cloud_record(cmod.update_archive_record_in_cloud, record_id=rec_id, fields={'barkod_seri_no': clean_new})
            except Exception as e:
                print(f"[CloudArchive] Async update hatasi: {e}")
            return True, "Barkod başarıyla güncellendi"

        conn.close()
        return True, "Barkod güncellendi"
    except Exception as e:
        return False, str(e)

def search_records(query='', start_date='', end_date='', limit=200):
    _ensure_cloud_synced()
    conn = get_connection()
    cursor = conn.cursor()
    sql = "SELECT * FROM barkod_kayitlari WHERE durum = 'Gecerli'"
    params = []

    if query:
        q = f"%{query.strip()}%"
        sql += " AND (kargo_barkodu LIKE ? OR barkod_seri_no LIKE ? OR siparis_no LIKE ? OR musteri_adi LIKE ? OR urun_adi LIKE ?)"
        params.extend([q, q, q, q, q])

    if start_date:
        sql += " AND kayit_gunu >= ?"
        params.append(start_date)

    if end_date:
        sql += " AND kayit_gunu <= ?"
        params.append(end_date)

    sql += " ORDER BY id DESC LIMIT ?"
    params.append(limit)

    cursor.execute(sql, params)
    rows = cursor.fetchall()
    conn.close()
    return [dict(r) for r in rows]

def get_stats():
    _ensure_cloud_synced()
    now = get_now_istanbul()
    bugun_str = now.strftime('%Y-%m-%d')

    try:
        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT COUNT(*) FROM barkod_kayitlari WHERE durum = 'Gecerli'")
        toplam_kayit = cursor.fetchone()[0]

        cursor.execute("SELECT COUNT(*) FROM barkod_kayitlari WHERE kayit_gunu = ? AND durum = 'Gecerli'", (bugun_str,))
        bugun_kayit = cursor.fetchone()[0]

        cursor.execute("SELECT COUNT(DISTINCT CASE WHEN kargo_barkodu != '' THEN kargo_barkodu ELSE siparis_no END) FROM barkod_kayitlari WHERE durum = 'Gecerli'")
        tekil_paket = cursor.fetchone()[0]

        cursor.execute("SELECT * FROM barkod_kayitlari WHERE durum = 'Gecerli' ORDER BY id DESC LIMIT 15")
        son_kayitlar = [dict(r) for r in cursor.fetchall()]

        conn.close()
    except Exception as e:
        print(f"Stats hatası: {e}")
        toplam_kayit = 0
        bugun_kayit = 0
        tekil_paket = 0
        son_kayitlar = []

    return {
        'toplam_kayit': toplam_kayit,
        'bugun_kayit': bugun_kayit,
        'tekil_siparis': tekil_paket,
        'bugun_tarih': bugun_str,
        'son_kayitlar': son_kayitlar
    }

def get_record_by_id(record_id):
    _ensure_cloud_synced()
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM barkod_kayitlari WHERE id = ?", (record_id,))
    row = cursor.fetchone()
    conn.close()
    return dict(row) if row else None

def export_records_to_excel(records, output_path=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kargo ve Barkod Kayıtları"

    headers = ["ID", "Kayıt Tarihi", "Kargo Barkodu", "Sipariş No", "Müşteri Adı", "Ürün Adı", "Ürün Seri No / Barkod", "Pazaryeri", "Okutma Yöntemi"]
    ws.append(headers)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    ws.row_dimensions[1].height = 28

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for row_idx, r in enumerate(records, start=2):
        ws.append([
            r.get('id', ''),
            r.get('kayit_tarihi', ''),
            r.get('kargo_barkodu', ''),
            r.get('siparis_no', ''),
            r.get('musteri_adi', ''),
            r.get('urun_adi', ''),
            r.get('barkod_seri_no', ''),
            r.get('pazaryeri', ''),
            r.get('okutma_yontemi', '')
        ])
        ws.row_dimensions[row_idx].height = 22
        is_even = (row_idx % 2 == 0)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill
            if col_idx in [1, 2, 3, 4, 8, 9]:
                cell.alignment = center_align
            elif col_idx in [3, 7]:
                cell.alignment = center_align
                cell.font = Font(name="Consolas", size=10.5, bold=True, color="1E3A8A")
            else:
                cell.alignment = left_align

    col_widths = {1: 8, 2: 20, 3: 22, 4: 18, 5: 25, 6: 35, 7: 25, 8: 15, 9: 16}
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    if output_path:
        wb.save(output_path)
    return wb


def save_active_orders_file(filename, orders):
    data_dir = get_base_data_dir()
    target = os.path.join(data_dir, 'active_orders.json')
    try:
        payload = {
            'filename': filename,
            'updated_at': get_now_istanbul().strftime('%Y-%m-%d %H:%M:%S'),
            'orders': orders
        }
        with open(target, 'w', encoding='utf-8') as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        return True
    except Exception as e:
        print(f'Aktif siparis kaydetme hatasi: {e}')
        return False

def get_active_orders_file():
    data_dir = get_base_data_dir()
    target = os.path.join(data_dir, 'active_orders.json')
    if os.path.exists(target):
        try:
            with open(target, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            print(f'Aktif siparis okuma hatasi: {e}')
    return None

def clear_active_orders_file():
    data_dir = get_base_data_dir()
    target = os.path.join(data_dir, 'active_orders.json')
    if os.path.exists(target):
        try:
            os.remove(target)
            return True
        except Exception:
            pass
    return False
