import os
import openpyxl
import csv

def clean_val(v):
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v)
    s = str(v).strip()
    if s.lower() in ('none', 'nan', 'null'):
        return ""
    if s.endswith('.0') and s[:-2].replace('-', '').isdigit():
        return s[:-2]
    return s

def normalize_header(s):
    if not s:
        return ""
    mapping = {
        'ç': 'c', 'ğ': 'g', 'ı': 'i', 'ö': 'o', 'ş': 's', 'ü': 'u',
        'Ç': 'c', 'Ğ': 'g', 'I': 'i', 'İ': 'i', 'Ö': 'o', 'Ş': 's', 'Ü': 'u'
    }
    s = str(s).lower()
    for k, v in mapping.items():
        s = s.replace(k, v)
    return "".join(c for c in s if c.isalnum())

def find_column_index(headers, possible_names):
    clean_headers = [normalize_header(h) for h in headers]
    for name in possible_names:
        n_clean = normalize_header(name)
        for idx, h_clean in enumerate(clean_headers):
            if h_clean == n_clean:
                return idx
        for idx, h_clean in enumerate(clean_headers):
            if n_clean in h_clean:
                return idx
    return None

def parse_order_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    rows = []

    if ext in ['.xlsx', '.xlsm']:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append([clean_val(c) for c in row])
    elif ext == '.csv':
        for enc in ['utf-8-sig', 'utf-8', 'windows-1254', 'iso-8859-9', 'latin-1']:
            try:
                with open(file_path, 'r', encoding=enc, errors='replace') as f:
                    reader = csv.reader(f)
                    file_rows = [row for row in reader if any(cell.strip() for cell in row)]
                    if file_rows:
                        rows = file_rows
                        break
            except Exception:
                continue
    elif ext == '.xls':
        try:
            import xlrd
            wb = xlrd.open_workbook(file_path)
            sheet = wb.sheet_by_index(0)
            for r in range(sheet.nrows):
                rows.append([clean_val(sheet.cell_value(r, c)) for c in range(sheet.ncols)])
        except Exception as e:
            print(f"XLS okuma hatası: {e}")
    else:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append([clean_val(c) for c in row])
        except Exception:
            pass

    if not rows:
        return []

    header_idx = -1
    for i in range(min(6, len(rows))):
        row_str = " ".join([normalize_header(c) for c in rows[i] if c is not None])
        if any(keyword in row_str for keyword in ['urun', 'siparis', 'adet', 'miktar', 'pazaryeri', 'barkod']):
            header_idx = i
            break

    if header_idx == -1:
        header_idx = 0

    headers = [str(c) if c is not None else "" for c in rows[header_idx]]
    clean_headers = [normalize_header(h) for h in headers]
    data_rows = rows[header_idx + 1:]

    col_kargo_barkod = find_column_index(headers, ['kargobarkod', 'kargobarkodu', 'kargotakipno', 'takipno', 'barkodkargo', 'kargono', 'gonderino', 'teslimatno', 'takipnumarasi'])
    col_kargo_firmasi = find_column_index(headers, ['kargofirmasi', 'kargo', 'carrier', 'kargosirketi'])
    col_siparis = find_column_index(headers, ['siparisno', 'siparisnumarasi', 'siparis', 'paketno', 'orderno', 'ordernumber', 'siparisid'])
    col_paket = find_column_index(headers, ['paketno', 'paketnumarasi', 'koli', 'kolino'])
    col_musteri = find_column_index(headers, [
        'teslimatadi', 'uyeadisoyadi', 'sevkmusteri', 'faturamusteri', 'musteriadi', 
        'alici', 'adsoyad', 'faturaadi', 'teslimatalicisi', 'aliciadi', 'musteri', 
        'aliciadsoyad', 'aliciunvan'
    ])
    col_urun = find_column_index(headers, ['urunadi', 'urun', 'productname', 'urunismi', 'aciklama', 'uruntanimi'])
    col_variant = find_column_index(headers, ['varyant', 'variant', 'secenek', 'ozellik'])
    col_adet = find_column_index(headers, ['miktar', 'adet', 'siparisadedi', 'quantity', 'qty', 'miktaradet'])
    col_pazar = find_column_index(headers, ['pazaryeri', 'pazaryeriadi', 'platform', 'magaza', 'kanal', 'kaynak', 'satici'])
    col_stok = find_column_index(headers, ['urunkodu', 'stokkodu', 'sku', 'modelkodu', 'urunid', 'stok'])
    col_gtin = find_column_index(headers, ['barkod', 'gtin', 'gtinbarkod', 'barcode', 'ean', 'ean13', 'gtin13'])
    col_telefon = find_column_index(headers, ['teslimattelefonu', 'musteritelefon', 'faturatelefonu', 'telefon', 'gsm', 'alicitelefon'])
    col_sehir = find_column_index(headers, ['teslimatsehiradi', 'faturasehiradi', 'sehir', 'il'])
    col_ilce = find_column_index(headers, ['teslimatilceadi', 'faturabolgeadi', 'ilce'])
    col_adres = find_column_index(headers, ['teslimatadresi', 'faturaadresi', 'adres'])

    # E-Ticaret / Web Sitesi (Pazar yerleri hariç doğrudan sipariş) tespiti
    is_web_export = (col_kargo_barkod is None) and any(k in clean_headers for k in [
        'odemetipi', 'uyeadisoyadi', 'teslimatadi', 'bankasiparisno', 'erpno', 'teslimatadresi', 'odemestipi'
    ])

    filename_lower = os.path.basename(file_path).lower()
    default_pazar = "Web Sitesi" if is_web_export else "Diğer"
    if "trendyol" in filename_lower:
        default_pazar = "Trendyol"
    elif "hepsiburada" in filename_lower or "hb" in filename_lower:
        default_pazar = "Hepsiburada"
    elif "cicek" in filename_lower:
        default_pazar = "Çiçeksepeti"
    elif "pazarama" in filename_lower:
        default_pazar = "Pazarama"
    elif "amazon" in filename_lower:
        default_pazar = "Amazon"
    elif "n11" in filename_lower:
        default_pazar = "N11"

    packages_map = {}
    for r_idx, r in enumerate(data_rows):
        if not r or not any(r):
            continue

        kargo_barkodu = clean_val(r[col_kargo_barkod]) if col_kargo_barkod is not None and col_kargo_barkod < len(r) else ""
        siparis_no = clean_val(r[col_siparis]) if col_siparis is not None and col_siparis < len(r) else ""
        paket_no = clean_val(r[col_paket]) if col_paket is not None and col_paket < len(r) else ""
        kargo_firmasi = clean_val(r[col_kargo_firmasi]) if col_kargo_firmasi is not None and col_kargo_firmasi < len(r) else ""
        musteri_adi = clean_val(r[col_musteri]) if col_musteri is not None and col_musteri < len(r) else "Belirtilmemiş"
        urun_adi = clean_val(r[col_urun]) if col_urun is not None and col_urun < len(r) else "Ürün"
        varyant = clean_val(r[col_variant]) if col_variant is not None and col_variant < len(r) else ""
        pazaryeri = clean_val(r[col_pazar]) if col_pazar is not None and col_pazar < len(r) else default_pazar
        stok_kod = clean_val(r[col_stok]) if col_stok is not None and col_stok < len(r) else ""
        gtin = clean_val(r[col_gtin]) if col_gtin is not None and col_gtin < len(r) else ""
        telefon = clean_val(r[col_telefon]) if col_telefon is not None and col_telefon < len(r) else ""
        sehir = clean_val(r[col_sehir]) if col_sehir is not None and col_sehir < len(r) else ""
        ilce = clean_val(r[col_ilce]) if col_ilce is not None and col_ilce < len(r) else ""
        adres = clean_val(r[col_adres]) if col_adres is not None and col_adres < len(r) else ""

        if varyant and varyant.lower() not in urun_adi.lower():
            urun_adi = f"{urun_adi} ({varyant})"

        adet = 1
        if col_adet is not None and col_adet < len(r) and r[col_adet] is not None:
            try:
                adet = int(float(str(r[col_adet]).replace(',', '.')))
                if adet < 1:
                    adet = 1
            except Exception:
                adet = 1

        # Kargo Barkodu yerine Sipariş No kuralı:
        # Kargo takip barkodu yoksa veya boşsa, Sipariş No doğrudan kargo barkodu olarak atanır.
        if kargo_barkodu:
            primary_kargo_code = kargo_barkodu
        elif siparis_no:
            primary_kargo_code = siparis_no
            kargo_barkodu = siparis_no
        elif paket_no:
            primary_kargo_code = paket_no
            kargo_barkodu = paket_no
        else:
            primary_kargo_code = f"SIPARIS-{r_idx+1}"
            kargo_barkodu = primary_kargo_code

        key = primary_kargo_code

        if key not in packages_map:
            packages_map[key] = {
                'kargo_barkodu': primary_kargo_code,
                'siparis_no': siparis_no,
                'paket_no': paket_no,
                'kargo_firmasi': kargo_firmasi,
                'musteri_adi': musteri_adi,
                'telefon': telefon,
                'sehir': sehir,
                'ilce': ilce,
                'adres': adres,
                'pazaryeri': pazaryeri,
                'is_web_order': is_web_export,
                'urunler': [],
                'tamamlandi': False
            }

        # Aynı pakette aynı ürün varsa (aynı stok kodu veya aynı ürün adı), adetleri birleştir
        existing_item = None
        for itm in packages_map[key]['urunler']:
            if stok_kod and itm.get('stok_kod'):
                if itm['stok_kod'].strip().lower() == stok_kod.strip().lower():
                    existing_item = itm
                    break
            elif itm.get('urun_adi') and itm['urun_adi'].strip().lower() == urun_adi.strip().lower():
                existing_item = itm
                break

        if existing_item:
            existing_item['adet'] += adet
        else:
            item_id = f"item_{len(packages_map[key]['urunler']) + 1}"
            packages_map[key]['urunler'].append({
                'item_id': item_id,
                'urun_adi': urun_adi,
                'stok_kod': stok_kod,
                'gtin': gtin,
                'adet': adet,
                'okunan_adet': 0,
                'seri_nolar': [],
                'tamamlandi': False
            })

    return list(packages_map.values())
